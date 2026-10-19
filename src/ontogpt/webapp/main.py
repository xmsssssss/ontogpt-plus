"""OntoGPT Webapp main function."""

import io
import json
import logging
import os
from io import StringIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import ontogpt.webapp.task_storage as task_storage

import docx
import openpyxl
import pymupdf  # PyMuPDF for PDF extraction
import uvicorn
from bs4 import BeautifulSoup
from fastapi import FastAPI, File, Form, Request, Response, UploadFile
from pydantic import BaseModel
from starlette.staticfiles import StaticFiles
from starlette.templating import Jinja2Templates

from ontogpt.cli import _get_templates
from ontogpt.engines.spires_engine import SPIRESEngine
from ontogpt.io.html_exporter import HTMLExporter
from ontogpt.io.json_wrapper import dump_minimal_json
from ontogpt.io.owl_exporter import OWLExporter
from ontogpt.io.rdf_exporter import RDFExporter
from ontogpt.io.template_loader import get_template_details
from ontogpt.io.yaml_wrapper import dump_minimal_yaml
import rdflib

this_path = Path(__file__).parent
static_dir = this_path / "static"
html_dir = this_path / "html"

# 中文模板解释映射表
TEMPLATE_ZH_DESCRIPTIONS: Dict[str, str] = {
    "power_grid": "【电力行业】电力系统设备、电网故障缺陷、二次保护动作与自适应参量全要素抽取 (IEC CIM / 国网标准)",
    "drug": "药物与疾病关联、作用机制、适应症及靶点抽取 (ChEBI / MONDO / HP)",
    "mendelian_disease": "单基因罕见遗传病、致病基因与表型特征 (MONDO / HP / HGNC)",
    "human_phenotype": "人类临床表型特征抽取与人类表型本体对齐 (HPO / HP)",
    "all_disease_grounding": "全量疾病名称抽取与 MONDO / 国际疾病本体对齐",
    "composite_disease": "复合型/多因素疾病综合概念抽取",
    "condition": "符合 FHIR 临床标准的病历状况与病症描述抽取",
    "diagnostic_procedure": "临床诊断检验流程与对应表型特征",
    "treatment": "医疗干预、治疗方案与处置手段 (MAXO)",
    "maxo": "MAXO 医学行动与治疗处置本体关联抽取",
    "phenopackets": "GA4GH Phenopackets 标准临床病例表型包抽取",
    "pathology": "临床病理学诊断与病理改变描述对齐",
    "dietician_notes": "临床营养师病历与营养干预记录抽取",
    "ctd": "化学物质与疾病因果关联关系 (BC5CDR 评测标准)",
    "ctd_ner": "化学物与疾病命名实体识别 (NER)",
    "recipe": "烹饪食谱：原料食材、配比、操作步骤与营养成分 (FOODON)",
    "food": "食品名称、食材成分与餐饮术语抽取 (FOODON)",
    "foodon_simple": "极简食品本体抽取 (FOODON)",
    "gene_extraction": "复杂上下文中的基因符号 (Gene Symbol) 识别与消歧",
    "go_term": "基因本体术语命名实体识别 (GO NER)",
    "go_terms": "基因本体 GO 术语与标准 ID 抽取 (GO:BP / GO:CC / GO:MF)",
    "go_terms_relational": "蛋白质参与的 GO 生物学过程与功能关系抽取",
    "go_simple": "极简基因本体术语抽取 (GO)",
    "gocam": "GO-CAM 基因本体因果活动模型抽取",
    "biological_process": "文献中的生物学过程与肾脏等器官机制抽取",
    "metabolic_process": "代谢过程与生化通路网络抽取",
    "cell_type": "细胞类型与解剖结构对齐抽取 (CL 细胞本体)",
    "emapa_simple": "小鼠发育解剖学本体抽取 (EMAPA)",
    "biomed_terminology": "通用生物医学术语标准抽取",
    "class-enrichment": "基因集富集分析结果的类富集模型",
    "mic": "微量营养素、生化反应通路及相关基因/疾病关系",
    "adverse_outcome_pathway": "有害结局路径 (AOP)：从分子起始事件到器官损伤",
    "alzrd": "阿尔茨海默病与痴呆症表型、实验指标与模式生物 (全文模式)",
    "alzrd_section": "阿尔茨海默病与痴呆症表型及实验指标 (分章节解析)",
    "ibd_literature": "炎症性肠病 (IBD) 文献核心知识抽取",
    "pet_breed": "猫狗等伴侣动物品种特征、表型与解剖部位 (VBO / UPHENO)",
    "vbo_names": "脊椎动物品种名称抽取 (VBO 脊椎动物品种本体)",
    "vbo_char": "动物品种名称及其表型外貌行为特征",
    "predator_prey": "生态学生物捕食与被捕食关系网络抽取",
    "biotic_interaction": "物种间生物互作关系 (共生/寄生/竞争)",
    "ecosim_methods": "EcoSIM 生态模拟方法抽取",
    "ecosim_simple": "极简 EcoSIM 生态抽取",
    "environmental-sample": "环境样本、采样地与环境特征 (ENVO)",
    "environmental-sample-ungrounded": "环境样本（仅保留单位，不对齐本体）",
    "environmental-metadata": "环境数据倡议 (EDI) 元数据分类抽取",
    "desiccation": "植物耐旱抗逆特征 (ChEBI / GO / PO / TO / PECO)",
    "metagenome": "环境宏基因组研究元数据与样本信息",
    "nmdc-schema-data": "国家微生物组数据协同中心 (NMDC) 标准环境元数据",
    "datasheet": "数据集说明文档规范元数据 (Datasheets for Datasets)",
    "data-sheets-schema": "数据集 Datasheet 规范模式",
    "table_arrays": "科研文献中的数值数据表格抽取 (表头+数值阵列)",
    "table_values": "PDF 文献中的行列数值表格抽取",
    "traits": "生物学表型性状特征抽取 (Traits)",
    "reaction": "生物化学反应方程式、底物与产物关系",
    "figure": "文献图表、插图标题与子部件元数据",
    "matrix_eval": "MATRIX 知识图谱项目的数据资源评估与评分",
    "miro": "MIRO 指南本体论文报告最小信息集规范",
    "mondo_simple": "极简 MONDO 疾病本体抽取",
    "onto_usage": "科学文献中本体的使用与引用声明抽取",
    "ontology_class": "通用知识本体概念类建模抽取",
    "ontology_issue": "GitHub 本体 Tracker 议题内容模型",
    "error_analysis": "数据校验错误汇总与聚合分析报告",
    "storms": "STORMS 肠道菌群研究规范清单抽取",
    "halo": "知识本体类假设与推理模型 (HALO)",
    "core": "核心顶层基础数据模型",
    "personinfo": "人员基础档案与联系信息",
}

DATAMODELS = {}

def refresh_datamodels() -> Dict[str, str]:
    global DATAMODELS
    all_templates = _get_templates()
    all_templates = dict(sorted(all_templates.items()))
    new_models = {}
    for template_id, (_name, description) in all_templates.items():
        zh_desc = TEMPLATE_ZH_DESCRIPTIONS.get(template_id)
        if zh_desc:
            new_models[template_id] = zh_desc
        else:
            new_models[template_id] = f"{description}"
    DATAMODELS = new_models
    return DATAMODELS

refresh_datamodels()

LLM_MODELS = [
    "custom",
    "gpt-4o",
    "gpt-4o-mini",
    "gpt-5.5",
    "gpt-5.4-mini",
    "anthropic/claude-sonnet-5",
    "anthropic/claude-haiku-4-5",
    "ollama/llama3",
    "ollama/qwen2.5",
    "ollama/gemma3",
]

LANGUAGE_SYSTEM_PROMPTS = {
    "auto": (
        "Output ONLY the requested fields in 'field_name: value' format. "
        "Do NOT include any conversational filler, markdown explanations, preambles, or notes."
    ),
    "zh_ground_en": (
        "You are an expert structured knowledge extraction engine. "
        "Extract the requested fields strictly in 'field_name: value1; value2' format. "
        "When extracting from Chinese text, translate recognized entities to their standard terminology. "
        "IMPORTANT RULES:\n"
        "1. Output ONLY key-value pairs matching the template's exact field names.\n"
        "2. Absolutely NO conversational preambles (e.g. 'Based on the text...', 'Here is the extraction...').\n"
        "3. Absolutely NO explanatory notes, comments, or disclaimers (e.g. '(Note:...')."
    ),
    "zh_native": (
        "你是一个顶级的工业与工程领域知识图谱抽取引擎。请从文本中严格按事实提取概念与关系：\n"
        "1. 严格忠实原文事实，绝不凭空臆造或脑补未提及的信息；\n"
        "2. 文本未明确提及的属性，坚决不要输出（绝对不要写'未明确指定'、'null'、'无'、'暂无'、'N/A'等无信息量占位符），直接忽略该行；\n"
        "3. 保持工业与工程专有名词、设备型号、技术规范标准代号（如 DL/T、GB/T 等）的完整性与专业性；\n"
        "4. 输出必须且只能包含请求的各个字段键值对（例如 equipment_components: 主变压器; 隔离开关），每行严格一个字段，不要把多个子槽位拼接在同一行；\n"
        "5. 绝不要输出任何礼貌用语、前言开场白、解释说明、Markdown附注或引用。"
    ),
}


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from uploaded files across various formats."""
    ext = Path(filename).suffix.lower()
    text = ""

    if ext in [".txt", ".md", ".json", ".yaml", ".yml", ".csv"]:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("gb18030")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="ignore")

    elif ext == ".pdf":
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        pages = []
        for page in doc:
            pages.append(page.get_text())
        doc.close()
        text = "\n\n".join(pages)

    elif ext == ".docx":
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                row_text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        text = "\n".join(paragraphs)

    elif ext in [".xlsx", ".xls"]:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheets_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines = [f"[工作表: {sheet_name}]"]
            for row in ws.iter_rows(values_only=True):
                non_empty = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if non_empty:
                    lines.append(", ".join(non_empty))
            sheets_text.append("\n".join(lines))
        text = "\n\n".join(sheets_text)

    elif ext in [".html", ".htm"]:
        soup = BeautifulSoup(file_bytes.decode("utf-8", errors="ignore"), "html.parser")
        for tag in soup(["script", "style", "nav", "footer"]):
            tag.decompose()
        text = soup.get_text(separator="\n", strip=True)

    else:
        # Fallback binary text decoding
        try:
            text = file_bytes.decode("utf-8")
        except Exception:
            text = file_bytes.decode("latin-1", errors="ignore")

    return text.strip()


class Query(BaseModel):
    text: str
    datamodel: str
    model: str


app = FastAPI(title="OntoGPT Web Studio")
static_dir.mkdir(parents=True, exist_ok=True)
app.mount("/static", StaticFiles(directory=str(static_dir)), name="static")
templates = Jinja2Templates(directory=str(html_dir))

import asyncio
import time
import uuid

html_exporter = HTMLExporter(output=None)
owl_exporter = OWLExporter()
rdf_exporter = RDFExporter()
engines: Dict[Tuple[str, str, Optional[str], str, str], SPIRESEngine] = {}

# 任务状态存储字典与事件订阅队列
# task_id -> { "status": "running"|"completed"|"failed"|"stopped", "step": int, "total": int, "percent": int, "message": str, "result": dict, "error": str, "engine": SPIRESEngine }
TASKS: Dict[str, Dict[str, Any]] = {}
TASK_QUEUES: Dict[str, List[asyncio.Queue]] = {}


def get_engine(
    datamodel: str,
    llm_model: str,
    api_base: Optional[str] = None,
    api_key: Optional[str] = None,
    custom_model: Optional[str] = None,
    lang_mode: str = "auto",
) -> SPIRESEngine:
    actual_model = (
        custom_model.strip() if (llm_model == "custom" and custom_model and custom_model.strip()) else llm_model
    )
    clean_api_base = api_base.strip() if (api_base and api_base.strip()) else None
    clean_api_key = api_key.strip() if (api_key and api_key.strip()) else ""
    system_message = LANGUAGE_SYSTEM_PROMPTS.get(lang_mode, "")

    model_provider = None
    if clean_api_base:
        model_provider = "openai"

    cache_key = (datamodel, actual_model, clean_api_base, clean_api_key, lang_mode)
    if cache_key not in engines:
        template_details = get_template_details(template=datamodel)
        engines[cache_key] = SPIRESEngine(
            model=actual_model,
            template_details=template_details,
            api_base=clean_api_base,
            api_key=clean_api_key,
            model_provider=model_provider,
            system_message=system_message,
        )
    return engines[cache_key]


@app.get("/")
def read_root(request: Request, task_id: Optional[str] = None):
    initial_text = None
    selected_datamodel = None
    selected_model = None
    if task_id:
        if task_id in TASKS:
            task_meta = TASKS[task_id].get("meta", {})
            initial_text = task_meta.get("input_text")
            selected_datamodel = task_meta.get("datamodel")
            selected_model = task_meta.get("llm_model")
        else:
            db_task = task_storage.get_task(task_id)
            if db_task:
                initial_text = db_task.get("input_text")
                selected_datamodel = db_task.get("datamodel")
                selected_model = db_task.get("llm_model")
    return templates.TemplateResponse(
        request=request,
        name="form.html",
        context={
            "request": request,
            "datamodels": refresh_datamodels(),
            "llm_models": LLM_MODELS,
            "initial_text": initial_text,
            "selected_datamodel": selected_datamodel,
            "selected_model": selected_model,
        },
    )


@app.post("/api/auto-generate-schema")
async def auto_generate_schema(
    text: str = Form(""),
    domain_hint: str = Form(""),
    schema_id: Optional[str] = Form(None),
    llm_model: str = Form("custom"),
    custom_model: Optional[str] = Form(None),
    api_base: Optional[str] = Form(None),
    api_key: Optional[str] = Form(None),
):
    """根据样例长文本或领域提示词，让大模型自主设计并编译 LinkML YAML 本体模板"""
    from ontogpt.clients.llm_client import LLMClient
    from ontogpt.io.schema_generator import (
        build_schema_prompt,
        compile_and_register_schema,
        extract_yaml_from_response,
        sanitize_schema_id,
    )

    actual_model = (
        custom_model.strip() if (llm_model == "custom" and custom_model and custom_model.strip()) else llm_model
    )
    clean_api_base = api_base.strip() if (api_base and api_base.strip()) else None
    clean_api_key = api_key.strip() if (api_key and api_key.strip()) else ""

    if not text.strip() and not domain_hint.strip():
        return {"success": False, "error": "请提供样例参考文本或领域名称/关键词！"}

    target_schema_id = sanitize_schema_id(schema_id or domain_hint or text[:20])
    if not target_schema_id or len(target_schema_id) < 3:
        target_schema_id = "custom_domain_schema"

    client = LLMClient(
        model=actual_model,
        api_base=clean_api_base,
        api_key=clean_api_key,
        custom_llm_provider="openai" if clean_api_base else None,
    )

    prompt = build_schema_prompt(
        sample_text=text or domain_hint,
        domain_hint=domain_hint or "工程与专业文献领域",
        schema_id=target_schema_id,
    )

    try:
        try:
            # 自动设计模板的输出较长，显式放宽输出上限，避免 YAML 结尾围栏被截断。
            raw_response = client.complete(prompt=prompt, max_tokens=8192)
        except SystemExit:
            # LLMClient 在不可恢复的 API 错误上会调用 sys.exit；这里兜住，
            # 避免一次模板生成请求把整个 Web 服务进程带走。
            return {
                "success": False,
                "error": "模型接口拒绝了本次请求（可能不支持该输出上限，或模型名与接口地址不匹配），请更换模型后重试。",
            }
        yaml_content = extract_yaml_from_response(raw_response)
        result = compile_and_register_schema(target_schema_id, yaml_content)
        refresh_datamodels()
        return {
            "success": True,
            "schema_id": target_schema_id,
            "title": result.get("title"),
            "classes": result.get("classes"),
            "root_class": result.get("root_class"),
            "yaml_content": result.get("yaml_content"),
            "message": f"成功自动设计并编译模板 [{target_schema_id}]！包含 {len(result.get('classes', []))} 个本体概念类。",
        }
    except Exception as e:
        import logging
        logging.exception("Auto schema generation failed")
        return {"success": False, "error": f"模板自动生成/编译失败: {str(e)}"}


@app.post("/api/parse-document")
async def parse_document(file: UploadFile = File(...)):
    """API endpoint to parse text from uploaded documents on the fly."""
    try:
        content = await file.read()
        extracted_text = extract_text_from_file(content, file.filename or "file.txt")
        return {
            "success": True,
            "filename": file.filename,
            "text": extracted_text,
            "char_count": len(extracted_text),
        }
    except Exception as e:
        return {"success": False, "error": f"解析失败: {str(e)}"}


from fastapi.responses import StreamingResponse

@app.post("/api/start-task")
async def start_task(
    datamodel: str = Form(...),
    text: Optional[str] = Form(None),
    llm_model: str = Form(...),
    custom_model: Optional[str] = Form(None),
    api_base: Optional[str] = Form(None),
    api_key: Optional[str] = Form(None),
    lang_mode: str = Form("zh_ground_en"),
    chunk_size: int = Form(6000),
    concurrency: int = Form(8),
    use_cache: bool = Form(True),
    auto_evolve: bool = Form(False),
    enable_neo4j_reuse: bool = Form(True),
    doc_file: Optional[UploadFile] = File(None),
):
    actual_model = (
        custom_model.strip() if (llm_model == "custom" and custom_model and custom_model.strip()) else llm_model
    )
    final_text = (text or "").strip()

    if doc_file and doc_file.filename:
        file_bytes = await doc_file.read()
        if len(file_bytes) > 0:
            doc_text = extract_text_from_file(file_bytes, doc_file.filename)
            if doc_text:
                final_text = doc_text

    if not final_text:
        return {"success": False, "error": "没有提供待抽取的文本内容！"}

    task_id = str(uuid.uuid4())
    engine = get_engine(
        datamodel=datamodel,
        llm_model=llm_model,
        api_base=api_base,
        api_key=api_key,
        custom_model=custom_model,
        lang_mode=lang_mode,
    )

    # 每次启动新任务时重置停止标志，并按需设置缓存策略
    engine.stop_requested = False
    engine.enable_neo4j_reuse = enable_neo4j_reuse
    if hasattr(engine, "client") and engine.client:
        engine.client.enable_cache = use_cache

    task_title = (
        doc_file.filename
        if (doc_file and doc_file.filename)
        else (final_text[:40].replace("\n", " ").strip() + "..." if len(final_text) > 40 else final_text)
    )
    task_storage.save_or_update_task(
        task_id=task_id,
        title=task_title,
        datamodel=datamodel,
        llm_model=actual_model,
        lang_mode=lang_mode,
        char_count=len(final_text),
        status="running",
        input_text=final_text,
    )

    TASKS[task_id] = {
        "status": "running",
        "step": 0,
        "total": 1,
        "percent": 5,
        "message": "任务初始化中...",
        "result": None,
        "error": None,
        "engine": engine,
        "start_time": time.time(),
        "meta": {
            "datamodel": datamodel,
            "llm_model": actual_model,
            "api_base": api_base or "",
            "lang_mode": lang_mode,
            "input_text": final_text,
        },
    }
    TASK_QUEUES[task_id] = []

    def run_extraction_sync():
        try:
            def on_progress(step: int, total: int, msg: str):
                TASKS[task_id]["step"] = step
                TASKS[task_id]["total"] = total
                # 进度计算：分块处理占 10% ~ 90%
                p = int(10 + (step / max(total, 1)) * 80)
                TASKS[task_id]["percent"] = min(p, 90)
                TASKS[task_id]["message"] = msg
                # 广播给 SSE 队列
                queues = TASK_QUEUES.get(task_id, [])
                for q in queues:
                    try:
                        q.put_nowait({"type": "progress", "percent": TASKS[task_id]["percent"], "message": msg})
                    except Exception:
                        pass

            engine.progress_callback = on_progress
            engine.max_workers = max(1, concurrency)

            if chunk_size > 0 and len(final_text) > chunk_size:
                engine.max_text_length = chunk_size
            else:
                engine.max_text_length = None

            on_progress(0, 1, "正在启动 SPIRES 提取引擎...")
            ann = engine.extract_from_text(final_text)

            if engine.stop_requested:
                raise InterruptedError("用户已手动停止任务")

            on_progress(1, 1, "正在生成知识图谱、OWL 公理与结构化报告...")
            output = StringIO()
            html_exporter.export(extraction_output=ann, output=output)
            inner_html = output.getvalue()
            yaml_output = dump_minimal_yaml(ann.model_dump())
            json_output = dump_minimal_json(ann.model_dump())
            owl_output = ""

            if engine.schemaview and ann.extracted_object:
                try:
                    owl_io = StringIO()
                    owl_exporter.export(ann, output=owl_io, schemaview=engine.schemaview, id_value=task_id)
                    owl_output = owl_io.getvalue()
                except Exception as owl_err:
                    logging.warning(f"OWLExporter failed: {owl_err}, fallback to RDFExporter")
                    try:
                        rdf_io = StringIO()
                        rdf_exporter.export(ann, output=rdf_io, schemaview=engine.schemaview)
                        rdf_ttl = rdf_io.getvalue()
                        if rdf_ttl.strip():
                            from ontogpt.io.owl_builder import clean_and_upgrade_owl_xml
                            g = rdflib.Graph()
                            g.parse(data=rdf_ttl, format="turtle")
                            raw_xml = g.serialize(format="pretty-xml", encoding="utf-8").decode("utf-8")
                            owl_output = clean_and_upgrade_owl_xml(raw_xml, schema_id=datamodel)
                    except Exception as rdf_err:
                        owl_output = f"# OWL export notice: {rdf_err}"

            # -------------------------------------------------------------
            # 模式自进化更新 (Auto-Evolve Schema)
            # -------------------------------------------------------------
            evolved_info = None
            if auto_evolve and ann and ann.extracted_object:
                raw_feats = getattr(ann.extracted_object, "auto_discovered_features", None) or []
                if raw_feats:
                    on_progress(1, 1, f"⚡ 正在执行模式自进化：对捕获的 {len(raw_feats)} 项动态专有特征进行知识建模并更新模板...")
                    try:
                        from ontogpt.io.schema_generator import auto_evolve_schema_with_features

                        discovered_features = []
                        for rf in raw_feats:
                            if isinstance(rf, dict):
                                discovered_features.append(rf)
                            else:
                                discovered_features.append(
                                    {"key": getattr(rf, "key", ""), "value": getattr(rf, "value", "")}
                                )

                        evolved_info = auto_evolve_schema_with_features(
                            schema_id=datamodel,
                            discovered_features=discovered_features,
                            client=engine.client,
                        )
                        if evolved_info.get("success") and evolved_info.get("evolved"):
                            # 清理内存中的引擎实例缓存，确保后续任务加载更新后的 Schema
                            keys_to_remove = [k for k in engines.keys() if k[0] == datamodel]
                            for k in keys_to_remove:
                                del engines[k]
                            refresh_datamodels()
                            logging.info(f"Schema {datamodel} successfully evolved with new features.")
                    except Exception as evo_err:
                        logging.warning(f"Auto-evolve schema failed: {evo_err}")

            finish_msg = "抽取与本体对齐全部完成！"
            if evolved_info and evolved_info.get("evolved"):
                finish_msg = f"抽取完成！模板 [{datamodel}] 已自动完成自进化更新并编译生效！"

            TASKS[task_id]["status"] = "completed"
            TASKS[task_id]["percent"] = 100
            TASKS[task_id]["message"] = finish_msg
            TASKS[task_id]["result"] = {
                "inner_html": inner_html,
                "yaml_output": yaml_output,
                "json_output": json_output,
                "owl_output": owl_output,
                "evolved_info": evolved_info,
            }
            task_storage.update_task_result(
                task_id=task_id,
                status="completed",
                result=TASKS[task_id]["result"],
            )
            queues = TASK_QUEUES.get(task_id, [])
            for q in queues:
                try:
                    q.put_nowait({"type": "completed", "percent": 100, "message": "完成！"})
                except Exception:
                    pass

        except InterruptedError:
            TASKS[task_id]["status"] = "stopped"
            TASKS[task_id]["message"] = "任务已被用户手动停止"
            task_storage.update_task_result(
                task_id=task_id,
                status="stopped",
                error="用户手动停止任务",
            )
            queues = TASK_QUEUES.get(task_id, [])
            for q in queues:
                try:
                    q.put_nowait({"type": "stopped", "message": "任务已停止"})
                except Exception:
                    pass
        except (Exception, SystemExit) as e:
            TASKS[task_id]["status"] = "failed"
            TASKS[task_id]["error"] = f"{type(e).__name__}: {str(e)}"
            TASKS[task_id]["message"] = f"任务异常失败: {str(e)}"
            task_storage.update_task_result(
                task_id=task_id,
                status="failed",
                error=TASKS[task_id]["error"],
            )
            queues = TASK_QUEUES.get(task_id, [])
            for q in queues:
                try:
                    q.put_nowait({"type": "failed", "error": TASKS[task_id]["error"]})
                except Exception:
                    pass

    # 异步线程池中执行
    asyncio.get_event_loop().run_in_executor(None, run_extraction_sync)
    return {"success": True, "task_id": task_id}


@app.post("/api/stop-task/{task_id}")
async def stop_task(task_id: str):
    """API endpoint to gracefully interrupt an in-progress extraction."""
    task = TASKS.get(task_id)
    if not task:
        return {"success": False, "error": "任务不存在或已过期"}

    if task["status"] == "running":
        engine = task.get("engine")
        if engine:
            engine.stop_requested = True
            # 若有运行中的线程池，立即强制撤销尚未启动的排队任务
            executor = getattr(engine, "_current_executor", None)
            if executor:
                try:
                    executor.shutdown(wait=False, cancel_futures=True)
                except Exception:
                    pass
        task["status"] = "stopped"
        task["message"] = "已发送停止指令，正在中断任务..."
        task_storage.update_task_result(task_id, status="stopped", error="用户手动停止任务")
        queues = TASK_QUEUES.get(task_id, [])
        for q in queues:
            try:
                q.put_nowait({"type": "stopped", "message": "任务已停止"})
            except Exception:
                pass
        return {"success": True, "message": "任务已停止"}
    return {"success": True, "message": f"当前状态: {task['status']}"}


@app.get("/api/cache-info")
def get_cache_info():
    """获取当前本地 LiteLLM 缓存状态"""
    db_path = ".litellm_cache/cache.db"
    count = 0
    size_mb = 0.0
    if os.path.exists(db_path):
        size_mb = round(os.path.getsize(db_path) / (1024 * 1024), 2)
        try:
            import sqlite3
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT count(*) FROM Cache;")
            count = cursor.fetchone()[0]
            conn.close()
        except Exception:
            pass
    return {"exists": os.path.exists(db_path), "count": count, "size_mb": size_mb}


@app.post("/api/clear-cache")
def clear_cache():
    """清空本地 LiteLLM 请求缓存"""
    db_path = ".litellm_cache/cache.db"
    try:
        if os.path.exists(db_path):
            import sqlite3
            conn = sqlite3.connect(db_path)
            conn.execute("DELETE FROM Cache;")
            conn.commit()
            conn.close()
        return {"success": True, "message": "已成功清空本地请求缓存！"}
    except Exception as e:
        return {"success": False, "error": f"清空缓存失败: {e}"}


@app.get("/api/task-status/{task_id}")
async def get_task_status(task_id: str):
    task = TASKS.get(task_id)
    if not task:
        return {"success": False, "error": "任务不存在"}
    return {
        "success": True,
        "status": task["status"],
        "step": task["step"],
        "total": task["total"],
        "percent": task["percent"],
        "message": task["message"],
        "error": task.get("error"),
    }


@app.get("/api/task-events/{task_id}")
async def task_events(task_id: str):
    """Server-Sent Events (SSE) 流式传输实时进度"""
    task = TASKS.get(task_id)
    if not task:
        return {"error": "Task not found"}

    queue = asyncio.Queue()
    if task_id in TASK_QUEUES:
        TASK_QUEUES[task_id].append(queue)

    async def event_generator():
        try:
            # 立即推送当前最新状态
            yield f"data: {json.dumps({'type': 'progress', 'percent': task['percent'], 'message': task['message'], 'status': task['status']}, ensure_ascii=False)}\n\n"
            while True:
                if task["status"] in ["completed", "failed", "stopped"]:
                    yield f"data: {json.dumps({'type': task['status'], 'percent': task['percent'], 'message': task['message'], 'error': task.get('error')}, ensure_ascii=False)}\n\n"
                    break
                try:
                    msg = await asyncio.wait_for(queue.get(), timeout=1.0)
                    yield f"data: {json.dumps(msg, ensure_ascii=False)}\n\n"
                    if msg.get("type") in ["completed", "failed", "stopped"]:
                        break
                except asyncio.TimeoutError:
                    # 心跳保持连接
                    yield f"data: {json.dumps({'type': 'ping'}, ensure_ascii=False)}\n\n"
        finally:
            if task_id in TASK_QUEUES and queue in TASK_QUEUES[task_id]:
                TASK_QUEUES[task_id].remove(queue)

    return StreamingResponse(event_generator(), media_type="text/event-stream")


@app.get("/task/{task_id}")
async def task_result_page(request: Request, task_id: str):
    """展示指定任务的最终抽取结果页面"""
    task = TASKS.get(task_id)
    if not task:
        db_task = task_storage.get_task(task_id)
        if db_task:
            task = {
                "id": db_task["id"],
                "status": db_task["status"],
                "meta": {
                    "datamodel": db_task["datamodel"],
                    "llm_model": db_task["llm_model"],
                    "lang_mode": db_task["lang_mode"],
                    "input_text": db_task.get("input_text", ""),
                },
                "result": db_task.get("result") or {},
                "error": db_task.get("error"),
            }
    if not task:
        return templates.TemplateResponse(
            request=request,
            name="form.html",
            context={"request": request, "datamodels": DATAMODELS, "llm_models": LLM_MODELS},
        )

    res = task.get("result") or {}
    meta = task.get("meta") or {}
    error_message = task.get("error")
    if task["status"] == "stopped":
        error_message = "抽取任务已由用户手动停止。"

    return templates.TemplateResponse(
        request=request,
        name="results.html",
        context={
            "request": request,
            "task_id": task_id,
            "datamodel": meta.get("datamodel", ""),
            "llm_model": meta.get("llm_model", ""),
            "api_base": meta.get("api_base", ""),
            "lang_mode": meta.get("lang_mode", ""),
            "input_text": meta.get("input_text", ""),
            "inner_html": res.get("inner_html", ""),
            "yaml_output": res.get("yaml_output", ""),
            "json_output": res.get("json_output", ""),
            "owl_output": res.get("owl_output", ""),
            "evolved_info": res.get("evolved_info"),
            "error_message": error_message,
        },
    )


@app.get("/history")
def history_page(request: Request):
    """展示任务历史记录与统计分析页面"""
    tasks = task_storage.list_tasks(limit=200)
    stats = task_storage.get_stats()
    return templates.TemplateResponse(
        request=request,
        name="history.html",
        context={"request": request, "tasks": tasks, "stats": stats},
    )


@app.get("/api/history")
def api_list_history(limit: int = 100, offset: int = 0):
    return {"success": True, "tasks": task_storage.list_tasks(limit, offset)}


@app.get("/api/history/stats")
def api_history_stats():
    return task_storage.get_stats()


@app.delete("/api/history/{task_id}")
def api_delete_task(task_id: str):
    if task_id in TASKS:
        del TASKS[task_id]
    success = task_storage.delete_task(task_id)
    return {"success": success}


@app.delete("/api/history")
def api_clear_all_history():
    TASKS.clear()
    success = task_storage.clear_all_tasks()
    return {"success": success}


@app.get("/api/export-merged-owl")
def api_export_merged_owl():
    """将所有已完成任务的 OWL 实体、关系与公理合并为一个全局 OWL 知识库文件 (Merged OWL Graph)"""
    merged_graph = rdflib.Graph()
    merged_graph.bind("owl", rdflib.OWL)
    merged_graph.bind("rdfs", rdflib.RDFS)
    merged_graph.bind("rdf", rdflib.RDF)

    tasks = task_storage.list_tasks(limit=500)
    for t in tasks:
        if t.get("status") == "completed" and t.get("has_result"):
            detail = task_storage.get_task(t["id"])
            if detail and detail.get("result"):
                owl_data = detail["result"].get("owl_output", "")
                if owl_data and "<rdf:RDF" in owl_data:
                    try:
                        merged_graph.parse(data=owl_data, format="xml")
                    except Exception as e:
                        logging.warning(f"Failed to merge task {t['id']} OWL: {e}")

    merged_raw_xml = merged_graph.serialize(format="pretty-xml", encoding="utf-8").decode("utf-8")
    from ontogpt.io.owl_builder import clean_and_upgrade_owl_xml
    merged_xml = clean_and_upgrade_owl_xml(merged_raw_xml, schema_id="merged_master_ontology")
    return Response(
        content=merged_xml,
        media_type="application/rdf+xml",
        headers={"Content-Disposition": "attachment; filename=merged_knowledge_base.owl"},
    )


@app.get("/explorer")
async def ontology_explorer_page(request: Request):
    """Ontology Explorer 页面：三栏式交互探索本体层级、属性公理与图谱结构"""
    tasks = task_storage.list_tasks(limit=100)
    completed_tasks = [t for t in tasks if t.get("status") == "completed" and t.get("has_result")]
    return templates.TemplateResponse(
        "explorer.html",
        {
            "request": request,
            "tasks": completed_tasks,
        },
    )


@app.get("/api/explorer/data")
def api_explorer_data(source: str = "merged", keyword: Optional[str] = None):
    """解析指定来源（全局融合或单个任务）的 OWL 为图形化交互数据结构，支持关键词关联子图筛选"""
    from ontogpt.webapp.ontology_inspector import parse_owl_to_explorer_model

    owl_xml = ""
    if source == "merged":
        merged_graph = rdflib.Graph()
        tasks = task_storage.list_tasks(limit=500)
        for t in tasks:
            if t.get("status") == "completed" and t.get("has_result"):
                detail = task_storage.get_task(t["id"])
                if detail and detail.get("result"):
                    owl_data = detail["result"].get("owl_output", "")
                    if owl_data and "<rdf:RDF" in owl_data:
                        try:
                            merged_graph.parse(data=owl_data, format="xml")
                        except Exception:
                            pass
        if len(merged_graph) > 0:
            owl_xml = merged_graph.serialize(format="pretty-xml", encoding="utf-8").decode("utf-8")
    else:
        detail = task_storage.get_task(source)
        if detail and detail.get("result"):
            owl_xml = detail["result"].get("owl_output", "")

    data = parse_owl_to_explorer_model(owl_xml, max_graph_nodes=100, search_keyword=keyword)
    return data


@app.get("/api/ontology/search")
def api_ontology_search(query: str, source: str = "merged", limit: int = 20):
    """
    语义知识查询开放接口 (Semantic Knowledge Search API)
    供外部系统、微服务、大模型 Agent (Tool / Function Call) 与前端直接调用的结构化本体知识服务。
    """
    from ontogpt.webapp.ontology_inspector import parse_owl_to_explorer_model

    clean_kw = query.strip()
    if not clean_kw:
        return {"query": query, "matched_count": 0, "results": []}

    owl_xml = ""
    if source == "merged":
        merged_graph = rdflib.Graph()
        tasks = task_storage.list_tasks(limit=500)
        for t in tasks:
            if t.get("status") == "completed" and t.get("has_result"):
                detail = task_storage.get_task(t["id"])
                if detail and detail.get("result"):
                    owl_data = detail["result"].get("owl_output", "")
                    if owl_data and "<rdf:RDF" in owl_data:
                        try:
                            merged_graph.parse(data=owl_data, format="xml")
                        except Exception:
                            pass
        if len(merged_graph) > 0:
            owl_xml = merged_graph.serialize(format="pretty-xml", encoding="utf-8").decode("utf-8")
    else:
        detail = task_storage.get_task(source)
        if detail and detail.get("result"):
            owl_xml = detail["result"].get("owl_output", "")

    explorer_data = parse_owl_to_explorer_model(owl_xml, max_graph_nodes=limit * 2, search_keyword=clean_kw)
    matched_inds = explorer_data.get("individuals", [])

    results = []
    classes_map = {c["uri"]: c for c in explorer_data.get("classes", [])}

    for ind in matched_inds[:limit]:
        type_uris = ind.get("types", [])
        primary_type_uri = type_uris[0] if type_uris else ""
        class_type_name = primary_type_uri.split("/")[-1] if primary_type_uri else "owl:Thing"

        # 过滤顶层聚合根节点（如 Report 容器），搜索实体时聚焦于具体的设备、规程、工艺与参数
        if "Report" in class_type_name:
            continue

        class_meta = classes_map.get(primary_type_uri, {})

        # 关系映射
        relationships = []
        for prop_name, targets in ind.get("object_properties", {}).items():
            for tgt in targets:
                relationships.append({
                    "predicate": prop_name,
                    "target_iri": tgt.get("uri"),
                    "target_name": tgt.get("label") or tgt.get("name"),
                })

        item_payload = {
            "entity": {
                "iri": ind.get("uri"),
                "name": ind.get("name"),
                "label": ind.get("label"),
                "class_type": primary_type_uri.split("/")[-1] if primary_type_uri else "owl:Thing",
                "class_hierarchy": [p.split("/")[-1] for p in class_meta.get("parents", [])] + [primary_type_uri.split("/")[-1]] if primary_type_uri else [],
            },
            "attributes": ind.get("data_properties", {}),
            "relationships": relationships,
            "characteristics": ["Functional" if any(len(v) == 1 for v in ind.get("data_properties", {}).values()) else "Standard"],
        }
        results.append(item_payload)

    return {
        "success": True,
        "query": clean_kw,
        "matched_count": len(results),
        "total_available": len(matched_inds),
        "results": results,
    }


@app.get("/api/neo4j/status")
def api_neo4j_status(
    uri: Optional[str] = None,
    user: Optional[str] = None,
    password: Optional[str] = None,
):
    """查询或测试 Neo4j 图数据库连接状态"""
    from ontogpt.clients.neo4j_service import neo4j_service, Neo4jKnowledgeService

    if uri or user or password:
        svc = Neo4jKnowledgeService(
            uri=uri or neo4j_service.uri,
            user=user or neo4j_service.user,
            password=password or neo4j_service.password,
        )
        res = svc.test_connection()
        svc.close()
        return res
    return neo4j_service.test_connection()


@app.post("/api/neo4j/sync")
async def api_neo4j_sync(request: Request):
    """将全局融合库或指定任务的本体知识同步至 Neo4j 图数据库"""
    from ontogpt.clients.neo4j_service import neo4j_service

    body = {}
    try:
        body = await request.json()
    except Exception:
        pass

    source = body.get("source", "merged")
    owl_xml = ""

    if source == "merged":
        merged_graph = rdflib.Graph()
        tasks = task_storage.list_tasks(limit=500)
        for t in tasks:
            if t.get("status") == "completed" and t.get("has_result"):
                detail = task_storage.get_task(t["id"])
                if detail and detail.get("result"):
                    owl_data = detail["result"].get("owl_output", "")
                    if owl_data and "<rdf:RDF" in owl_data:
                        try:
                            merged_graph.parse(data=owl_data, format="xml")
                        except Exception:
                            pass
        if len(merged_graph) > 0:
            owl_xml = merged_graph.serialize(format="pretty-xml", encoding="utf-8").decode("utf-8")
    else:
        detail = task_storage.get_task(source)
        if detail and detail.get("result"):
            owl_xml = detail["result"].get("owl_output", "")

    if not owl_xml:
        return {"success": False, "error": "没有可供同步的本体数据"}

    sync_res = neo4j_service.sync_owl_to_neo4j(owl_xml)
    return sync_res


@app.post("/")
async def form_post(
    request: Request,
    datamodel: str = Form(...),
    text: Optional[str] = Form(None),
    llm_model: str = Form(...),
    custom_model: Optional[str] = Form(None),
    api_base: Optional[str] = Form(None),
    api_key: Optional[str] = Form(None),
    lang_mode: str = Form("zh_ground_en"),
    chunk_size: int = Form(3500),
    doc_file: Optional[UploadFile] = File(None),
):
    actual_model = (
        custom_model.strip() if (llm_model == "custom" and custom_model and custom_model.strip()) else llm_model
    )
    final_text = (text or "").strip()

    # If user uploaded a document directly via form submit
    if doc_file and doc_file.filename:
        file_bytes = await doc_file.read()
        if len(file_bytes) > 0:
            doc_text = extract_text_from_file(file_bytes, doc_file.filename)
            if doc_text:
                final_text = doc_text

    error_message = None
    inner_html = ""
    yaml_output = ""
    json_output = ""
    owl_output = ""

    if not final_text:
        error_message = "没有提供待抽取的文本内容。请粘贴文本或上传包含文字的文档！"
    else:
        try:
            engine = get_engine(
                datamodel=datamodel,
                llm_model=llm_model,
                api_base=api_base,
                api_key=api_key,
                custom_model=custom_model,
                lang_mode=lang_mode,
            )

            # 长文档智能自适应分块保护：
            # 若文本超过单次请求设定长度，自动切块分批提示 LLM 并递归合并去重
            if chunk_size > 0 and len(final_text) > chunk_size:
                engine.max_text_length = chunk_size
            else:
                engine.max_text_length = None

            ann = engine.extract_from_text(final_text)
            output = StringIO()
            html_exporter.export(extraction_output=ann, output=output)
            inner_html = output.getvalue()
            yaml_output = dump_minimal_yaml(ann.model_dump())
            json_output = dump_minimal_json(ann.model_dump())

            if engine.schemaview and ann.extracted_object:
                try:
                    owl_io = StringIO()
                    owl_exporter.export(ann, output=owl_io, schemaview=engine.schemaview)
                    owl_output = owl_io.getvalue()
                except Exception as owl_err:
                    logging.warning(f"OWLExporter failed: {owl_err}, fallback to RDFExporter")
                    try:
                        rdf_io = StringIO()
                        rdf_exporter.export(ann, output=rdf_io, schemaview=engine.schemaview)
                        rdf_ttl = rdf_io.getvalue()
                        if rdf_ttl.strip():
                            from ontogpt.io.owl_builder import clean_and_upgrade_owl_xml
                            g = rdflib.Graph()
                            g.parse(data=rdf_ttl, format="turtle")
                            raw_xml = g.serialize(format="pretty-xml", encoding="utf-8").decode("utf-8")
                            owl_output = clean_and_upgrade_owl_xml(raw_xml, schema_id=datamodel)
                    except Exception as rdf_err:
                        owl_output = f"# OWL export notice: {rdf_err}"
        except BaseException as e:
            error_message = f"{type(e).__name__}: {e}"

    return templates.TemplateResponse(
        request=request,
        name="results.html",
        context={
            "request": request,
            "datamodel": datamodel,
            "llm_model": actual_model,
            "api_base": api_base or "",
            "lang_mode": lang_mode,
            "input_text": final_text,
            "inner_html": inner_html,
            "yaml_output": yaml_output,
            "json_output": json_output,
            "owl_output": owl_output,
            "error_message": error_message,
        },
    )


def start():
    """Launch with `uv run web-ontogpt` or `web-ontogpt`."""
    uvicorn.run(
        "ontogpt.webapp.main:app",
        host="127.0.0.1",
        port=8000,
        reload=True,
        reload_dirs=["src"],
        reload_excludes=[".venv", ".git", "__pycache__", "task_history.db", "*.db", "*.log"],
    )
